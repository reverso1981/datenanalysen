from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Rechenmodell_Seesauna_Attersee.xlsx")

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="1F4E79")
YELLOW = PatternFill("solid", fgColor="FFFF00")
SUMFILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Border(top=Side(style="thin"))

EUR = '#,##0;(#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
DEC = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
YRS = '0.0" J."'

wb = Workbook()


def hdr(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill = HDR, HDRFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, bold=False):
    c = ws[ref]
    c.value = value
    c.font = Font(name=F, bold=bold, color=font.color) if bold else font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


# ---------------------------------------------------------------- Annahmen
A = wb.active
A.title = "Annahmen"
A["A1"] = "Seesauna Attersee – Annahmen (Rahmenkonzept, Stand Oktober 2026)"
A["A1"].font = TITLE
A["A2"] = ("Legende: blaue Zahlen = Eingaben (änderbar) · schwarz = Formeln · grün = Verweise auf andere Blätter · "
           "gelb hinterlegt = Schlüsselannahmen. Szenario in C4 wählen (1 = pessimistisch, 2 = Basis, 3 = optimistisch).")
A["A2"].font = Font(name=F, italic=True)
A["A2"].alignment = Alignment(wrap_text=True)
A.merge_cells("A2:H2")
A.row_dimensions[2].height = 32
A["B4"] = "Aktives Szenario (1/2/3)"
A["B4"].font = BOLD
put(A, "C4", 2, BLUE, fill=YELLOW)
A["D4"] = '=INDEX({"pessimistisch","Basis","optimistisch"},1,C4)'
A["D4"].font = BOLD

hdr_row = 6
for i, t in enumerate(["Key", "Annahme", "pessimistisch", "Basis", "optimistisch", "aktiv", "Einheit", "Quelle / Hinweis"], 1):
    c = A.cell(row=hdr_row, column=i, value=t)
    c.font, c.fill = HDR, HDRFILL
for col, w in zip("ABCDEFGH", [14, 46, 14, 14, 14, 14, 12, 70]):
    A.column_dimensions[col].width = w

R = {}  # key -> row in Annahmen
assumptions = [
    ("§", "Kapazität & Besuch"),
    ("tage_hs", "Öffnungstage Hauptsaison (Okt–Apr, abzgl. 7 Revisionstage)", None, 205, None, "Tage", "Ganzjahresbetrieb lt. Grilling; 212 Tage Okt–Apr minus Revision", False),
    ("tage_ns", "Öffnungstage Nebensaison (Mai–Sep, abzgl. 7 Tage)", None, 146, None, "Tage", "Im Sommer reduzierte Zeiten / Kombi mit Strandbad", False),
    ("g_hs", "Ø Besucher pro Tag Hauptsaison", 140, 180, 215, "Besucher", "Annahme; ~1,5 Belegungswechsel bei 120 Gleichzeitig-Plätzen", True),
    ("g_ns", "Ø Besucher pro Tag Nebensaison", 50, 70, 90, "Besucher", "Annahme; Sommer schwach, Seezugang + Strandbad-Kombi als Treiber", True),
    ("kap", "Max. Besucher pro Tag (Spitzentag)", None, 300, None, "Besucher", "120 gleichzeitig × ~2,5 Wechsel", False),
    ("§", "Preise (brutto, inkl. USt)"),
    ("p_tag", "Tageskarte", 36, 39, 42, "€", "Zielband 35–42 € lt. Grilling (gehobene Mitte)", True),
    ("p_3h", "3-Stunden-Karte", 23, 25, 27, "€", "lt. Grilling ca. 25 €", False),
    ("a_tag", "Anteil Tageskarten an Besuchen", None, 0.55, None, "%", "Annahme", False),
    ("rabatt", "Ø Rabatt (Abos, Einheimischentarif, Hotelkontingente)", 0.15, 0.12, 0.10, "%", "Annahme; Einheimischentarif als Auflage der Gemeinden", False),
    ("ust_e", "USt-Satz Eintritt", None, 0.13, None, "%", "Annahme ermäßigter Satz für Bäder/Saunen – steuerlich prüfen", False),
    ("§", "Gastronomie & Zusatzerlöse"),
    ("g_kopf", "Gastro-Umsatz pro Besucher (brutto)", 11, 12.5, 14, "€", "Branchenüblich 8–15 € (Bistro, kleine Karte)", True),
    ("ust_g", "Ø USt-Satz Gastro (Mix 10 %/20 %)", None, 0.15, None, "%", "Speisen 10 %, Getränke 20 %", False),
    ("we_g", "Wareneinsatzquote Gastro", None, 0.32, None, "%", "Branchenwert Bistro 28–35 %", False),
    ("kiosk", "Strandbad-Kiosk Sommer (Umsatz netto)", 45000, 60000, 75000, "€/Jahr", "Annahme; Betreiber übernimmt Kiosk Mai–Sep", False),
    ("we_k", "Wareneinsatzquote Kiosk", None, 0.40, None, "%", "Annahme", False),
    ("m_anz", "Massagebehandlungen pro Jahr (2 Räume)", 3000, 4000, 5000, "Stück", "Annahme; freie Masseur:innen", False),
    ("m_preis", "Ø Preis Massage (brutto)", None, 75, None, "€", "Annahme", False),
    ("ust_m", "USt-Satz Massage", None, 0.20, None, "%", "Normalsteuersatz", False),
    ("m_prov", "Provision / Raummiete an Betreiber", None, 0.25, None, "%", "Annahme; Anteil am Netto-Behandlungsumsatz", False),
    ("ev_anz", "Events & Exklusivbuchungen pro Jahr", 25, 40, 55, "Stück", "Vollmond-/Themenaufgüsse, Firmen-Buyouts", False),
    ("ev_erl", "Ø Zusatzerlös je Event (netto)", None, 2500, None, "€", "Annahme", False),
    ("§", "Betriebskosten Betreiber"),
    ("fte", "Personal (Vollzeitäquivalente)", 21, 20, 20, "VZÄ", "Basis = Blatt Stellenplan; pessimistisch +1 VZÄ", True),
    ("k_fte", "Personalkosten je VZÄ (inkl. Lohnnebenkosten)", None, 52000, None, "€/Jahr", "Ø lt. Blatt Stellenplan (inkl. Zuschläge, Dienstgeberanteil, Aushilfen)", False),
    ("energie", "Energiekosten (Seethermie-WP + PV + Holz)", 220000, 180000, 160000, "€/Jahr", "Annahme; ohne Seethermie/PV eher 300–400 Tsd. €", True),
    ("wasser", "Wasser / Abwasser", None, 40000, None, "€/Jahr", "Annahme", False),
    ("waesche", "Wäsche, Reinigungs- & Verbrauchsmaterial je Besuch", None, 2.2, None, "€", "Annahme", False),
    ("mkt", "Marketing (in % vom Umsatz)", None, 0.04, None, "%", "Annahme", False),
    ("wart", "Laufende Wartung Betreiber (in % vom Umsatz)", None, 0.02, None, "%", "Kleinreparaturen; große Instandhaltung = Gemeinde-GmbH", False),
    ("verw", "Verwaltung, Versicherung, IT, Beratung", None, 85000, None, "€/Jahr", "Annahme", False),
    ("§", "Pachtvertrag (Gemeinde-GmbH → Betreiber)"),
    ("p_satz", "Umsatzpacht (in % vom Nettoumsatz)", None, 0.10, None, "%", "Zielband 8–12 % lt. Grilling", True),
    ("p_min", "Mindestpacht", None, 150000, None, "€/Jahr", "Annahme; Absicherung der Gemeinden", False),
    ("ziel_m", "Ziel-EBITDA-Marge Betreiber (Prüfgröße tragbare Pacht)", None, 0.08, None, "%", "Untergrenze, damit ein privater Betreiber Risiko übernimmt", False),
    ("§", "Gemeinde-Errichtungs-GmbH"),
    ("rl", "Instandhaltungsrücklage (in % der Bauwerkskosten KG 2–4)", None, 0.008, None, "%", "Faustwert 0,8–1,2 % p. a.", False),
    ("g_verw", "Verwaltung Errichtungs-GmbH", None, 25000, None, "€/Jahr", "Buchhaltung, Prüfung, Versicherung Gebäude", False),
    ("komm", "Kommunalsteuer (in % der Lohnsumme)", None, 0.03, None, "%", "KommStG § 9; fließt an Standortgemeinde", False),
    ("lohnant", "Anteil Bruttolohn an Personalkosten", None, 0.76, None, "%", "Annahme; Rest = Dienstgeberabgaben", False),
    ("§", "Entwicklung"),
    ("start", "Eröffnungsjahr", None, 2029, None, "Jahr", "Annahme: 2027 Planung/Genehmigung, 2028 Bau", False),
    ("hl1", "Hochlauf Jahr 1 (in % des Normaljahres)", 0.70, 0.75, 0.85, "%", "Annahme", False),
    ("hl2", "Hochlauf Jahr 2", 0.85, 0.90, 0.95, "%", "Annahme", False),
    ("inf_p", "Preissteigerung p. a.", None, 0.025, None, "%", "Annahme", False),
    ("inf_k", "Kostensteigerung p. a.", None, 0.03, None, "%", "Annahme", False),
]

row = hdr_row + 1
for a in assumptions:
    if a[0] == "§":
        A.cell(row=row, column=2, value=a[1]).font = BOLD
        row += 1
        continue
    key, label, pess, basis, opt, unit, src, key_flag = a
    R[key] = row
    A.cell(row=row, column=1, value=key).font = Font(name=F, color="808080", size=8)
    A.cell(row=row, column=2, value=label).font = BLACK
    fmt = PCT if unit == "%" else (DEC if isinstance(basis, float) and basis < 100 else NUM)
    if unit == "Jahr":
        fmt = "0"
    d = A.cell(row=row, column=4, value=basis)
    d.font, d.number_format = BLUE, fmt
    for col, v in ((3, pess), (5, opt)):
        c = A.cell(row=row, column=col)
        if v is None:
            c.value, c.font = f"=D{row}", BLACK
        else:
            c.value, c.font = v, BLUE
        c.number_format = fmt
    f = A.cell(row=row, column=6, value=f"=INDEX(C{row}:E{row},1,$C$4)")
    f.number_format, f.font = fmt, BOLD
    A.cell(row=row, column=7, value=unit).font = BLACK
    A.cell(row=row, column=8, value=src).font = Font(name=F, size=9, italic=True)
    if key_flag:
        for col in (3, 4, 5):
            A.cell(row=row, column=col).fill = YELLOW
    row += 1
A.freeze_panes = "C7"


def a(key, col):
    """Reference to an assumption in scenario column col (C/D/E/F)."""
    return f"Annahmen!${col}${R[key]}"


SC = ["C", "D", "E"]  # scenario columns, identical layout on Umsatz/Betriebskosten/Szenarien

# ---------------------------------------------------------------- Raumprogramm
RP = wb.create_sheet("Raumprogramm")
RP["A1"] = "Raumprogramm (Nutzflächen)"
RP["A1"].font = TITLE
hdr(RP, 3, ["Bereich", "Nutzfläche (m²)", "Plätze", "Erläuterung"], [44, 16, 10, 70])
rooms = [
    ("Eingang, Kassa, Shop-Ecke", 60, None, "Zutrittssystem mit Chip-Armband, Gutscheinverkauf"),
    ("Garderobe & Umkleiden (Spinde)", 140, 180, "180 Spinde für 120 Gleichzeitig-Gäste + Wechsel"),
    ("Sanitär & Duschen", 100, None, "Getrennte Duschzonen, Saunavorduschen"),
    ("Panorama-Aufgusssauna 90 °C", 55, 35, "Event-Aufgüsse, Glasfront Seeblick"),
    ("Bio-Sauna 60 °C", 25, 12, "Sanfte Sauna, Farblicht"),
    ("Dampfbad", 22, 10, "Inkl. Technikraum Dampferzeuger"),
    ("Holzofen-Saunahaus am Steg (außen)", 30, 15, "Separates Häuschen, direkter Seezugang"),
    ("Abkühlbereich innen (Tauchbecken, Kneipp, Schwallduschen)", 70, None, "Tauchbecken ca. 6 m²"),
    ("Ruheraum (Stille)", 200, 50, "50 Liegen à ~4 m², Seeblick"),
    ("Kommunikationsraum (Liegen, Gespräche erlaubt)", 200, 50, "50 Liegen, Kamin, Bibliothek"),
    ("Bistro (Gastraum + Theke)", 90, 40, "Im Bademantel; Terrasse saisonal zusätzlich"),
    ("Bistro-Küche (Regenerierküche) & Lager", 50, None, "Kleine Karte, 3–4 warme Gerichte"),
    ("Massage-/Behandlungsräume", 40, 2, "2 Räume à 20 m²"),
    ("Personal, Büro, Sozialräume", 55, None, ""),
    ("Haustechnik, Wäsche, Lager", 150, None, "Wärmepumpe, Lüftung, Wasseraufbereitung"),
    ("Verkehrsflächen", 140, None, "Gänge, Treppen"),
]
r0 = 4
for i, (n, m2, pl, t) in enumerate(rooms):
    r = r0 + i
    RP.cell(row=r, column=1, value=n).font = BLACK
    put(RP, f"B{r}", m2, BLUE, NUM)
    if pl is not None:
        put(RP, f"C{r}", pl, BLUE, NUM)
    RP.cell(row=r, column=4, value=t).font = Font(name=F, size=9, italic=True)
r_last = r0 + len(rooms) - 1
r_nf = r_last + 1
put(RP, f"A{r_nf}", "Summe Nutzfläche (NF)", bold=True)
put(RP, f"B{r_nf}", f"=SUM(B{r0}:B{r_last})", BLACK, NUM, SUMFILL, bold=True)
r_fak = r_nf + 1
put(RP, f"A{r_fak}", "Faktor BGF / NF")
put(RP, f"B{r_fak}", 1.18, BLUE, DEC)
RP[f"D{r_fak}"] = "Annahme kompakter Holzbau (Wände, Schächte)"
RP[f"D{r_fak}"].font = Font(name=F, size=9, italic=True)
r_bgf = r_fak + 1
put(RP, f"A{r_bgf}", "Brutto-Grundfläche (BGF)", bold=True)
put(RP, f"B{r_bgf}", f"=ROUND(B{r_nf}*B{r_fak},-1)", BLACK, NUM, SUMFILL, bold=True)
r_sauna = r_bgf + 2
put(RP, f"A{r_sauna}", "Saunaplätze gesamt (Kontrolle)")
put(RP, f"C{r_sauna}", f"=C7+C8+C9+C10", BLACK, NUM)
put(RP, f"A{r_sauna+1}", "Liegeplätze Ruhe + Kommunikation")
put(RP, f"C{r_sauna+1}", f"=C12+C13", BLACK, NUM)
BGF = f"Raumprogramm!$B${r_bgf}"

# ---------------------------------------------------------------- Investition
I = wb.create_sheet("Investition")
I["A1"] = "Einmalige Investition (netto, exkl. USt – Vorsteuerabzug über Vermietung der Gemeinde-GmbH)"
I["A1"].font = TITLE
hdr(I, 3, ["KG", "Position", "Menge", "Einheit", "Einheitspreis (€)", "Summe (€)", "Hinweis"], [6, 52, 10, 9, 16, 16, 62])
inv = [
    ("1", "Aufschließung", [
        ("Erschließung, Parkplatzanpassung, Leitungen", 1, "psch", 200000, "Strandbad-Infrastruktur teilweise vorhanden"),
        ("Abbruch / Anpassung Bestand Strandbad", 1, "psch", 100000, ""),
    ]),
    ("2", "Bauwerk – Rohbau", [
        ("Holzbau-Rohbau inkl. Fundierung, Dach", "BGF", "m²", 1100, "Regionaler Holzbau (CLT/Holzrahmen)"),
    ]),
    ("3", "Bauwerk – Technik", [
        ("Haustechnik HLS, Lüftung, Elektro", "BGF", "m²", 900, "Hohe Lüftungs-/Feuchteanforderungen"),
        ("Seewasser-Wärmepumpe inkl. Entnahme/Rückgabe", 1, "psch", 450000, "Wasserrechtliche Bewilligung nötig"),
        ("Photovoltaik", 150, "kWp", 1100, "Dachfläche ca. 800 m²"),
        ("Wasseraufbereitung Tauchbecken/Whirlpool", 1, "psch", 120000, "ÖNORM M 6215 / Bäderhygienegesetz"),
    ]),
    ("4", "Bauwerk – Ausbau", [
        ("Innenausbau, Glasfronten, Oberflächen", "BGF", "m²", 1100, "Gehobener Standard, große Seeverglasung"),
        ("Panorama-Aufgusssauna (35 Plätze)", 1, "Stk", 260000, "Inkl. Ofen, Steuerung, Glasfront"),
        ("Bio-Sauna (12 Plätze)", 1, "Stk", 80000, ""),
        ("Dampfbad (10 Plätze)", 1, "Stk", 130000, "Inkl. Dampferzeuger, Fliesen/Mineralwerkstoff"),
        ("Holzofen-Saunahaus am Steg (15 Plätze)", 1, "Stk", 220000, "Massivholz, holzbefeuert, Funkenschutz"),
        ("Tauchbecken innen", 1, "Stk", 80000, ""),
        ("Außen-Whirlpool (ca. 12 Personen)", 1, "Stk", 150000, "Beheizt über Wärmepumpe"),
    ]),
    ("5", "Einrichtung", [
        ("Liegen (Ruhe, Kommunikation, Garten)", 130, "Stk", 900, "Hochwertige Relaxliegen"),
        ("Gastro-Küche & Theke", 1, "psch", 180000, "Regenerierküche, Kühlung, Kaffeetechnik"),
        ("Möblierung, Deko, Leitsystem", 1, "psch", 200000, ""),
        ("Kassen-/Zutrittssystem, Spinde, IT", 1, "psch", 200000, "Chip-Armband, Online-Ticketing"),
    ]),
    ("6", "Außenanlagen", [
        ("Saunagarten, Sichtschutz zum Strandbad, Wege", 1, "psch", 300000, "Textilfrei → Sichtschutz zwingend"),
        ("Steg / Seezugang (ganzjährig, Leiter, Beleuchtung)", 1, "psch", 250000, "Wasserrecht, Schifffahrt, Naturschutz"),
    ]),
]
r = 4
kg_sum_rows = {}
for kg, title, items in inv:
    I.cell(row=r, column=1, value=kg).font = BOLD
    I.cell(row=r, column=2, value=title).font = BOLD
    r += 1
    first = r
    for name, qty, unit, ep, note in items:
        I.cell(row=r, column=2, value=name).font = BLACK
        if qty == "BGF":
            put(I, f"C{r}", f"={BGF}", GREEN, NUM)
        else:
            put(I, f"C{r}", qty, BLUE, NUM)
        I.cell(row=r, column=4, value=unit).font = BLACK
        put(I, f"E{r}", ep, BLUE, EUR)
        put(I, f"F{r}", f"=C{r}*E{r}", BLACK, EUR)
        I.cell(row=r, column=7, value=note).font = Font(name=F, size=9, italic=True)
        r += 1
    put(I, f"B{r}", f"Summe KG {kg}", bold=True)
    put(I, f"F{r}", f"=SUM(F{first}:F{r-1})", BLACK, EUR, SUMFILL, bold=True)
    kg_sum_rows[kg] = r
    r += 2

r_kg16 = r
put(I, f"B{r}", "Zwischensumme KG 1–6", bold=True)
put(I, f"F{r}", "=" + "+".join(f"F{kg_sum_rows[k]}" for k in "123456"), BLACK, EUR, SUMFILL, bold=True)
r += 1
r_kg7 = r
put(I, f"A{r}", "7", bold=True)
put(I, f"B{r}", "Planungsleistungen (Architektur, Fachplanung, ÖBA)")
put(I, f"E{r}", 0.15, BLUE, PCT)
put(I, f"F{r}", f"=F{r_kg16}*E{r}", BLACK, EUR)
I[f"G{r}"] = "Anteil an KG 1–6; Richtwert 14–18 %"
r += 1
r_kg8 = r
put(I, f"A{r}", "8", bold=True)
put(I, f"B{r}", "Nebenleistungen (Gutachten, Genehmigungen, Projektsteuerung)")
put(I, f"E{r}", 0.03, BLUE, PCT)
put(I, f"F{r}", f"=F{r_kg16}*E{r}", BLACK, EUR)
I[f"G{r}"] = "Wasserrecht, Naturschutz, Bodengutachten, Bauherrenvertretung"
r += 1
r_kg9 = r
put(I, f"A{r}", "9", bold=True)
put(I, f"B{r}", "Reserve (Unvorhergesehenes, Indexierung)")
put(I, f"E{r}", 0.12, BLUE, PCT, YELLOW)
put(I, f"F{r}", f"=(F{r_kg16}+F{r_kg7}+F{r_kg8})*E{r}", BLACK, EUR)
I[f"G{r}"] = "Konzeptphase: 10–15 % üblich"
r += 1
r_total = r
put(I, f"B{r}", "GESAMTINVESTITION netto", bold=True)
put(I, f"F{r}", f"=F{r_kg16}+F{r_kg7}+F{r_kg8}+F{r_kg9}", BLACK, EUR, YELLOW, bold=True)
r += 1
put(I, f"B{r}", "je m² BGF")
put(I, f"F{r}", f"=F{r_total}/{BGF}", BLACK, EUR)
r += 1
put(I, f"B{r}", "je Gleichzeitig-Gast (120)")
put(I, f"F{r}", f"=F{r_total}/120", BLACK, EUR)
r += 1
r_bauwerk = r
put(I, f"B{r}", "Bauwerkskosten KG 2–4 (Basis Instandhaltungsrücklage)")
put(I, f"F{r}", f"=F{kg_sum_rows['2']}+F{kg_sum_rows['3']}+F{kg_sum_rows['4']}", BLACK, EUR)
r += 1
put(I, f"B{r}", "Bauwerkskosten KG 2–4 je m² BGF")
put(I, f"F{r}", f"=F{r_bauwerk}/{BGF}", BLACK, EUR)

r += 2
put(I, f"B{r}", "Förderungen (Annahmen – mit Fördergebern abzuklären)", bold=True)
r += 1
# find WP and PV rows
wp_row = pv_row = None
for rr in range(4, r_kg16):
    v = I.cell(row=rr, column=2).value
    if v and v.startswith("Seewasser"):
        wp_row = rr
    if v == "Photovoltaik":
        pv_row = rr
r_f1 = r
put(I, f"B{r}", "UFI/KPC Umweltförderung (Seethermie + PV)")
put(I, f"E{r}", 0.30, BLUE, PCT)
put(I, f"F{r}", f"=(F{wp_row}+F{pv_row})*E{r}", BLACK, EUR)
I[f"G{r}"] = "Quote auf förderfähige Kosten Energieanlage; Annahme"
r += 1
put(I, f"B{r}", "Land OÖ (Tourismusinfrastruktur / Bedarfszuweisungen)")
put(I, f"E{r}", 0.20, BLUE, PCT, YELLOW)
put(I, f"F{r}", f"=F{r_total}*E{r}", BLACK, EUR)
I[f"G{r}"] = "Größter Hebel – politisch mit Land OÖ zu verhandeln; Annahme"
r += 1
put(I, f"B{r}", "LEADER (Region Vöckla-Ager / Attersee)")
put(I, f"F{r}", 150000, BLUE, EUR)
I[f"G{r}"] = "Pauschale Annahme; LEADER-Förderungen sind gedeckelt"
r_f3 = r
r += 1
r_foerd = r
put(I, f"B{r}", "Summe Förderungen", bold=True)
put(I, f"F{r}", f"=SUM(F{r_f1}:F{r_f3})", BLACK, EUR, SUMFILL, bold=True)
r += 2
r_gem = r
put(I, f"B{r}", "VON GEMEINDEN + TVB ZU TRAGEN (einmalig, netto)", bold=True)
put(I, f"F{r}", f"=F{r_total}-F{r_foerd}", BLACK, EUR, YELLOW, bold=True)
r += 1
put(I, f"B{r}", "Förderquote")
put(I, f"F{r}", f"=F{r_foerd}/F{r_total}", BLACK, PCT)
for rr in range(4, r + 1):
    I.cell(row=rr, column=7).font = Font(name=F, size=9, italic=True)
I.freeze_panes = "A4"

INV_TOTAL = f"Investition!$F${r_total}"
INV_GEM = f"Investition!$F${r_gem}"
INV_BAU = f"Investition!$F${r_bauwerk}"
INV_FOERD = f"Investition!$F${r_foerd}"

# ---------------------------------------------------------------- Umsatz (Normaljahr, je Szenario)
U = wb.create_sheet("Umsatz")
U["A1"] = "Umsatz Betreiber – Normaljahr (ab Jahr 3, Preisbasis Eröffnung, netto)"
U["A1"].font = TITLE
hdr(U, 3, ["Position", "Einheit", "pessimistisch", "Basis", "optimistisch", "aktiv"], [50, 10, 16, 16, 16, 16])
U_ROWS = {}
urows = [
    ("bes_hs", "Besuche Hauptsaison", "Besuche", lambda c: f"={a('tage_hs', c)}*{a('g_hs', c)}", NUM),
    ("bes_ns", "Besuche Nebensaison", "Besuche", lambda c: f"={a('tage_ns', c)}*{a('g_ns', c)}", NUM),
    ("bes", "Besuche gesamt", "Besuche", lambda c: f"={c}{{bes_hs}}+{c}{{bes_ns}}", NUM),
    ("ausl", "Auslastung Hauptsaison (Ø / Spitzentag)", "%", lambda c: f"={a('g_hs', c)}/{a('kap', c)}", PCT),
    ("ein_b", "Ø Eintritt je Besuch brutto", "€", lambda c: f"=({a('a_tag', c)}*{a('p_tag', c)}+(1-{a('a_tag', c)})*{a('p_3h', c)})*(1-{a('rabatt', c)})", DEC),
    ("ein_n", "Ø Eintritt je Besuch netto", "€", lambda c: f"={c}{{ein_b}}/(1+{a('ust_e', c)})", DEC),
    (None, "", None, None, None),
    ("u_ein", "Eintrittserlöse", "€", lambda c: f"={c}{{bes}}*{c}{{ein_n}}", EUR),
    ("u_gas", "Gastronomie (Bistro)", "€", lambda c: f"={c}{{bes}}*{a('g_kopf', c)}/(1+{a('ust_g', c)})", EUR),
    ("u_kio", "Strandbad-Kiosk (Sommer)", "€", lambda c: f"={a('kiosk', c)}", EUR),
    ("u_mas", "Massage (Provision/Raummiete)", "€", lambda c: f"={a('m_anz', c)}*{a('m_preis', c)}/(1+{a('ust_m', c)})*{a('m_prov', c)}", EUR),
    ("u_evt", "Events & Exklusivbuchungen", "€", lambda c: f"={a('ev_anz', c)}*{a('ev_erl', c)}", EUR),
    ("u_sum", "UMSATZ GESAMT (netto)", "€", lambda c: f"=SUM({c}{{u_ein}}:{c}{{u_evt}})", EUR),
    ("u_kopf", "Umsatz je Besuch", "€", lambda c: f"={c}{{u_sum}}/{c}{{bes}}", DEC),
]
r = 4
for key, label, unit, fn, fmt in urows:
    if key:
        U_ROWS[key] = r
    r += 1
r = 4
for key, label, unit, fn, fmt in urows:
    if key is None:
        r += 1
        continue
    U.cell(row=r, column=1, value=label).font = BOLD if key in ("bes", "u_sum") else BLACK
    U.cell(row=r, column=2, value=unit).font = BLACK
    for c in SC:
        f = fn(c).format(**U_ROWS)
        pure_link = f.count("!") == 1 and not any(op in f[1:] for op in "*/+-")
        cell = put(U, f"{c}{r}", f, GREEN if pure_link else BLACK, fmt)
        if key in ("bes", "u_sum"):
            cell.fill = SUMFILL
            cell.font = BOLD
    put(U, f"F{r}", f"=INDEX(C{r}:E{r},1,Annahmen!$C$4)", BLACK, fmt)
    if key in ("bes", "u_sum"):
        U[f"F{r}"].font = BOLD
        U[f"F{r}"].fill = SUMFILL
    r += 1
U.freeze_panes = "C4"


def u(key, c):
    return f"Umsatz!{c}{U_ROWS[key]}"


# ---------------------------------------------------------------- Betriebskosten (Normaljahr, je Szenario)
B = wb.create_sheet("Betriebskosten")
B["A1"] = "Betriebskosten & Ergebnis Betreiber – Normaljahr (netto)"
B["A1"].font = TITLE
hdr(B, 3, ["Position", "Art", "pessimistisch", "Basis", "optimistisch", "aktiv"], [50, 10, 16, 16, 16, 16])
brows = [
    ("k_weg", "Wareneinsatz Gastro", "variabel", lambda c: f"={u('u_gas', c)}*{a('we_g', c)}"),
    ("k_wek", "Wareneinsatz Kiosk", "variabel", lambda c: f"={u('u_kio', c)}*{a('we_k', c)}"),
    ("k_pers", "Personal", "fix", lambda c: f"={a('fte', c)}*{a('k_fte', c)}"),
    ("k_ene", "Energie", "fix", lambda c: f"={a('energie', c)}"),
    ("k_was", "Wasser / Abwasser", "fix", lambda c: f"={a('wasser', c)}"),
    ("k_wae", "Wäsche, Reinigung, Verbrauchsmaterial", "variabel", lambda c: f"={u('bes', c)}*{a('waesche', c)}"),
    ("k_mkt", "Marketing", "variabel", lambda c: f"={u('u_sum', c)}*{a('mkt', c)}"),
    ("k_wart", "Laufende Wartung", "variabel", lambda c: f"={u('u_sum', c)}*{a('wart', c)}"),
    ("k_verw", "Verwaltung, Versicherung, IT", "fix", lambda c: f"={a('verw', c)}"),
    ("k_sum", "BETRIEBSKOSTEN vor Pacht", "", lambda c: f"=SUM({c}{{k_weg}}:{c}{{k_verw}})"),
    ("ebitda_v", "Ergebnis vor Pacht (EBITDAR)", "", lambda c: f"={u('u_sum', c)}-{c}{{k_sum}}"),
    ("pacht", "Pacht an Gemeinde-GmbH", "", lambda c: f"=MAX({a('p_min', c)},{a('p_satz', c)}*{u('u_sum', c)})"),
    ("ebitda", "BETREIBERERGEBNIS nach Pacht (EBITDA)", "", lambda c: f"={c}{{ebitda_v}}-{c}{{pacht}}"),
    (None, "", None, None),
    ("marge", "EBITDA-Marge", "%", lambda c: f"={c}{{ebitda}}/{u('u_sum', c)}"),
    ("q_pers", "Personalkostenquote", "%", lambda c: f"={c}{{k_pers}}/{u('u_sum', c)}"),
    ("q_ene", "Energiekostenquote", "%", lambda c: f"={c}{{k_ene}}/{u('u_sum', c)}"),
    ("q_pacht", "Pachtquote", "%", lambda c: f"={c}{{pacht}}/{u('u_sum', c)}"),
    ("fix", "Fixkosten (Personal, Energie, Wasser, Verwaltung)", "", lambda c: f"={c}{{k_pers}}+{c}{{k_ene}}+{c}{{k_was}}+{c}{{k_verw}}"),
    ("db", "Deckungsbeitrag je Besuch (nach Umsatzpacht)", "€", lambda c: f"=({u('u_sum', c)}*(1-{a('p_satz', c)})-{c}{{k_weg}}-{c}{{k_wek}}-{c}{{k_wae}}-{c}{{k_mkt}}-{c}{{k_wart}})/{u('bes', c)}"),
    ("be", "Break-even Besuche pro Jahr (Betreiber)", "Besuche", lambda c: f"={c}{{fix}}/{c}{{db}}"),
    ("be_q", "Break-even in % der geplanten Besuche", "%", lambda c: f"={c}{{be}}/{u('bes', c)}"),
]
B_ROWS = {}
r = 4
for k, *_ in brows:
    if k:
        B_ROWS[k] = r
    r += 1
r = 4
for k, label, art, fn in brows:
    if k is None:
        r += 1
        continue
    big = k in ("k_sum", "ebitda_v", "ebitda", "pacht")
    B.cell(row=r, column=1, value=label).font = BOLD if big else BLACK
    B.cell(row=r, column=2, value=art).font = BLACK
    fmt = PCT if art == "%" else (DEC if art == "€" else EUR)
    for c in SC:
        f = fn(c).format(**B_ROWS)
        cell = put(B, f"{c}{r}", f, BLACK, fmt)
        if big:
            cell.font, cell.fill = BOLD, SUMFILL
    put(B, f"F{r}", f"=INDEX(C{r}:E{r},1,Annahmen!$C$4)", BLACK, fmt)
    if big:
        B[f"F{r}"].font, B[f"F{r}"].fill = BOLD, SUMFILL
    r += 1
B.freeze_panes = "C4"


def b(key, c):
    return f"Betriebskosten!{c}{B_ROWS[key]}"


# ---------------------------------------------------------------- Stellenplan
SP = wb.create_sheet("Stellenplan")
SP["A1"] = "Stellenbesetzungsplan Betreiber (Normaljahr, Gehälter = Annahmen)"
SP["A1"].font = TITLE
SP["A2"] = ("Gehälter orientiert am KV Hotel- und Gastgewerbe (Arbeiter/Angestellte) mit marktüblicher Überzahlung im "
            "Salzkammergut – vor Verwendung mit dem aktuell gültigen KV abgleichen.")
SP["A2"].font = Font(name=F, italic=True, color="C00000")
SP.merge_cells("A2:J2")
for col, w in zip("ABCDEFGHIJ", [40, 13, 9, 15, 15, 14, 15, 16, 14, 52]):
    SP.column_dimensions[col].width = w
sp_in = [
    ("Öffnungsstunden pro Tag (10–22 Uhr)", 12, NUM, BLUE, "Annahme"),
    ("Vor-/Nachbereitung pro Tag (h)", 1, NUM, BLUE, "Aufheizen, Reinigung, Kassenabschluss"),
    ("Betriebstage pro Jahr", f"={a('tage_hs','D')}+{a('tage_ns','D')}", NUM, GREEN, "aus Annahmen"),
    ("Betriebsstunden pro Jahr", "=(B4+B5)*B6", NUM, BLACK, ""),
    ("Wochenarbeitszeit (h)", 40, NUM, BLUE, "KV Hotel- und Gastgewerbe"),
    ("Urlaub (Wochen)", 5, NUM, BLUE, "Urlaubsgesetz"),
    ("Feiertage (Tage)", 13, NUM, BLUE, "gesetzliche Feiertage AT"),
    ("Krankenstand (Tage, Ø)", 10, NUM, BLUE, "Annahme"),
    ("Netto-Jahresstunden je VZÄ", "=B8*52-B9*B8-(B10+B11)*B8/5", NUM, BLACK, "verfügbare Arbeitszeit"),
    ("VZÄ je durchgehend besetzter Position", "=B7/B12", DEC, BLACK, "Faustregel: 1 Dauerposition ≈ 2,7 VZÄ"),
    ("Gehälter pro Jahr", 14, NUM, BLUE, "inkl. Urlaubs- und Weihnachtsgeld"),
    ("Zuschläge Sonn-/Feiertag, Abend (in % der Bruttogehälter)", 0.06, PCT, BLUE, "7-Tage-Betrieb; Annahme"),
    ("Dienstgeberanteil (SV, DB, DZ, KommSt, MVK)", 0.29, PCT, BLUE, "Richtwert AT ca. 28–30 %"),
    ("Aushilfen / Geringfügige (pauschal, inkl. Abgaben)", 30000, EUR, BLUE, "Spitzen: Wochenende, Events, Sommer-Kiosk"),
]
for i, (lab, v, fmt, fnt, note) in enumerate(sp_in):
    rr = 4 + i
    SP.cell(row=rr, column=1, value=lab).font = BLACK
    put(SP, f"B{rr}", v, fnt, fmt)
    SP.cell(row=rr, column=10, value=note).font = Font(name=F, size=9, italic=True)
hdr(SP, 19, ["Funktion", "Bereich", "VZÄ", "Brutto/Monat (€)", "Jahresbrutto (€)", "Zuschläge (€)",
             "DG-Anteil (€)", "Personalkosten (€)", "Ø Besetzung je Betriebsstunde", "Aufgaben / Einstufung"])
jobs = [
    ("Betriebsleitung", "Leitung", 1.0, 4300, "Gesamtverantwortung, Personal, Budget, Marketing; Tagdienst + Rufbereitschaft"),
    ("Stv. Leitung / Gästeservice-Leitung", "Leitung", 1.0, 3200, "Dienstplanung, Kassa/Rezeption, Hotelkooperationen, Events"),
    ("Haustechnik / Bademeister:in", "Technik", 1.0, 3100, "Wärmepumpe, Wasseraufbereitung, Saunaöfen, Holzofen; Bäderhygiene"),
    ("Saunameister:in – Teamleitung", "Sauna", 1.0, 2850, "Aufgussplan, Event-Aufgüsse, Einschulung"),
    ("Saunameister:innen", "Sauna", 4.0, 2500, "Aufgüsse, Aufsicht Saunabereich & Steg, Erste Hilfe"),
    ("Rezeption / Kassa", "Gästeservice", 4.0, 2200, "Check-in, Verkauf, Gutscheine, Telefon/Online"),
    ("Reinigung", "Service", 3.0, 2050, "Laufende Reinigung Nass-/Ruhebereiche, Wäschelogistik"),
    ("Koch / Köchin (Regenerierküche)", "Gastro", 2.0, 2450, "Kleine Karte, 3–4 warme Gerichte"),
    ("Service / Theke Bistro", "Gastro", 3.0, 2150, "Theke, Bademantel-Bistro, im Sommer Kiosk"),
]
j0 = 20
for i, (fn_, area, vz, br, note) in enumerate(jobs):
    rr = j0 + i
    SP.cell(row=rr, column=1, value=fn_).font = BLACK
    SP.cell(row=rr, column=2, value=area).font = BLACK
    put(SP, f"C{rr}", vz, BLUE, '0.0')
    put(SP, f"D{rr}", br, BLUE, EUR)
    put(SP, f"E{rr}", f"=C{rr}*D{rr}*$B$14", BLACK, EUR)
    put(SP, f"F{rr}", f"=E{rr}*$B$15", BLACK, EUR)
    put(SP, f"G{rr}", f"=(E{rr}+F{rr})*$B$16", BLACK, EUR)
    put(SP, f"H{rr}", f"=E{rr}+F{rr}+G{rr}", BLACK, EUR)
    put(SP, f"I{rr}", f"=C{rr}*$B$12/$B$7", BLACK, DEC)
    SP.cell(row=rr, column=10, value=note).font = Font(name=F, size=9, italic=True)
jl = j0 + len(jobs) - 1
js = jl + 1
put(SP, f"A{js}", "Summe Stammpersonal", bold=True)
put(SP, f"C{js}", f"=SUM(C{j0}:C{jl})", BLACK, '0.0', SUMFILL, bold=True)
for col in "EFGHI":
    put(SP, f"{col}{js}", f"=SUM({col}{j0}:{col}{jl})", BLACK, DEC if col == "I" else EUR, SUMFILL, bold=True)
put(SP, f"A{js+1}", "Aushilfen / Geringfügige")
put(SP, f"H{js+1}", "=B17", BLACK, EUR)
SP_TOT = js + 2
put(SP, f"A{SP_TOT}", "PERSONALKOSTEN GESAMT", bold=True)
put(SP, f"H{SP_TOT}", f"=H{js}+H{js+1}", BLACK, EUR, YELLOW, bold=True)
SP_AVG = SP_TOT + 1
put(SP, f"A{SP_AVG}", "Ø Personalkosten je VZÄ (→ Annahmen)")
put(SP, f"H{SP_AVG}", f"=H{SP_TOT}/C{js}", BLACK, EUR)
put(SP, f"A{SP_AVG+1}", "Personalkosten in % vom Umsatz (aktives Szenario)")
put(SP, f"H{SP_AVG+1}", f"=H{SP_TOT}/{u('u_sum','F')}", BLACK, PCT)
SP[f"A{SP_AVG+3}"] = ("Lesart Spalte I: Wie viele Personen dieser Funktion im Schnitt je Betriebsstunde im Haus sind. "
                      "Summe ≈ 7–8 Personen gleichzeitig; Leitung/Technik arbeiten überwiegend im Tagdienst.")
SP[f"A{SP_AVG+3}"].font = Font(name=F, size=9, italic=True)
SP.freeze_panes = "A20"
# Annahmen auf Stellenplan verlinken
rf, rk = R["fte"], R["k_fte"]
put(A, f"D{rf}", f"=Stellenplan!C{js}", GREEN, NUM, YELLOW)
put(A, f"C{rf}", f"=D{rf}+1", BLACK, NUM, YELLOW)
put(A, f"D{rk}", f"=Stellenplan!H{SP_AVG}", GREEN, NUM)

# ---------------------------------------------------------------- GuV 10 Jahre (aktives Szenario)
G = wb.create_sheet("GuV_10J")
G["A1"] = "Betreiber-GuV & Pachtfluss – 10 Jahre (aktives Szenario, nominal)"
G["A1"].font = TITLE
G["A2"] = '="Szenario: "&Annahmen!D4'
G["A2"].font = BOLD
G.column_dimensions["A"].width = 46
for j in range(10):
    G.column_dimensions[L(2 + j)].width = 13
hdr(G, 4, ["€"] + [f"J{j+1}" for j in range(10)])
YC = [L(2 + j) for j in range(10)]
grows = [
    ("jahr", "Kalenderjahr", lambda j, c: f'=TEXT({a("start", "F")}+{j},"0")', "@"),
    ("hl", "Hochlauf-Faktor", lambda j, c: f"={a('hl1','F')}" if j == 0 else (f"={a('hl2','F')}" if j == 1 else "=1"), PCT),
    ("pi", "Preisindex", lambda j, c: f"=(1+{a('inf_p','F')})^{j}", DEC),
    ("ki", "Kostenindex", lambda j, c: f"=(1+{a('inf_k','F')})^{j}", DEC),
    ("bes", "Besuche", lambda j, c: f"={u('bes','F')}*{c}{{hl}}", NUM),
    ("ums", "Umsatz netto", lambda j, c: f"={u('u_sum','F')}*{c}{{hl}}*{c}{{pi}}", EUR),
    ("var", "Variable Kosten (Wareneinsatz, Wäsche, Mkt, Wartung)", lambda j, c: f"=({b('k_weg','F')}+{b('k_wek','F')}+{b('k_wae','F')}+{b('k_mkt','F')}+{b('k_wart','F')})*{c}{{hl}}*{c}{{ki}}", EUR),
    ("fix", "Fixkosten (Personal, Energie, Wasser, Verwaltung)", lambda j, c: f"={b('fix','F')}*{c}{{ki}}", EUR),
    ("ebr", "Ergebnis vor Pacht", lambda j, c: f"={c}{{ums}}-{c}{{var}}-{c}{{fix}}", EUR),
    ("pac", "Pacht", lambda j, c: f"=MAX({a('p_min','F')}*{c}{{ki}},{a('p_satz','F')}*{c}{{ums}})", EUR),
    ("ebd", "Betreiberergebnis (EBITDA)", lambda j, c: f"={c}{{ebr}}-{c}{{pac}}", EUR),
    ("ebk", "Betreiberergebnis kumuliert", lambda j, c: f"={c}{{ebd}}" if j == 0 else f"={YC[j-1]}{{ebk}}+{c}{{ebd}}", EUR),
    (None, "Gemeinde-Errichtungs-GmbH", None, None),
    ("g_pac", "Pachteinnahme", lambda j, c: f"={c}{{pac}}", EUR),
    ("g_rl", "Instandhaltungsrücklage", lambda j, c: f"=-{INV_BAU}*{a('rl','F')}*{c}{{ki}}", EUR),
    ("g_vw", "Verwaltung GmbH", lambda j, c: f"=-{a('g_verw','F')}*{c}{{ki}}", EUR),
    ("g_ue", "Überschuss für Rückführung an Gemeinden", lambda j, c: f"=SUM({c}{{g_pac}}:{c}{{g_vw}})", EUR),
    ("g_kum", "Rückfluss kumuliert", lambda j, c: f"={c}{{g_ue}}" if j == 0 else f"={YC[j-1]}{{g_kum}}+{c}{{g_ue}}", EUR),
    ("g_ko", "Kommunalsteuer an Standortgemeinde (zusätzlich)", lambda j, c: f"={b('k_pers','F')}*{a('lohnant','F')}*{a('komm','F')}*{c}{{ki}}", EUR),
]
G_ROWS = {}
r = 5
for k, *_ in grows:
    if k:
        G_ROWS[k] = r
    r += 1
r = 5
for k, label, fn, fmt in grows:
    G.cell(row=r, column=1, value=label).font = BOLD if k in (None, "ums", "ebd", "g_ue") else BLACK
    if k is None:
        r += 1
        continue
    for j, c in enumerate(YC):
        cell = put(G, f"{c}{r}", fn(j, c).format(**G_ROWS), BLACK, fmt)
        if k in ("ums", "ebd", "g_ue"):
            cell.font, cell.fill = BOLD, SUMFILL
    r += 1
r += 1
put(G, f"A{r}", "Summe Rückfluss 10 Jahre", bold=True)
put(G, f"B{r}", f"=SUM(B{G_ROWS['g_ue']}:K{G_ROWS['g_ue']})", BLACK, EUR, bold=True)
r += 1
put(G, f"A{r}", "Rückflussquote 10 J. auf Gemeinde-Investition")
put(G, f"B{r}", f"=B{r-1}/{INV_GEM}", BLACK, PCT)
G.freeze_panes = "B5"

# ---------------------------------------------------------------- Pacht/Gemeinden
P = wb.create_sheet("Gemeinden")
P["A1"] = "Aufteilung der einmaligen Investition auf Gemeinden & TVB"
P["A1"].font = TITLE
P["A2"] = ("Einwohner- und Nächtigungszahlen sind GROBE PLATZHALTER und vor Verwendung mit Statistik Austria "
           "(Bevölkerung 1.1., Tourismusstatistik) bzw. Land OÖ zu ersetzen. Standortgemeinde mit 'x' markieren.")
P["A2"].font = Font(name=F, italic=True, color="C00000")
P["A2"].alignment = Alignment(wrap_text=True)
P.merge_cells("A2:H2")
P.row_dimensions[2].height = 32
put(P, "A4", "Zu verteilende Investition (netto nach Förderung)")
put(P, "D4", f"={INV_GEM}", GREEN, EUR)
put(P, "A5", "Anteil Tourismusverband Attersee-Attergau")
put(P, "D5", 0.10, BLUE, PCT, YELLOW)
put(P, "A6", "Bonus-Anteil Standortgemeinde")
put(P, "D6", 0.15, BLUE, PCT, YELLOW)
put(P, "A7", "Gewichtung Einwohner (Rest = Nächtigungen)")
put(P, "D7", 0.50, BLUE, PCT, YELLOW)
put(P, "E5", "Begründung: TVB profitiert über Nächtigungen/Angebot")
put(P, "E6", "Begründung: Kommunalsteuer, Frequenz, Aufwertung Strandbad")
hdr(P, 9, ["Gemeinde", "Einwohner", "Nächtigungen/Jahr", "Standort (x)", "Anteil", "Einmalbetrag (€)", "Ø Pachtrückfluss p. a. (€)", "€ je Einwohner"],
    [34, 13, 18, 12, 10, 18, 22, 14])
P.column_dimensions["A"].width = 46
gem = [
    ("Seewalchen am Attersee", 5700, 75000, ""),
    ("Schörfling am Attersee", 3500, 45000, ""),
    ("Attersee am Attersee", 1650, 120000, "x"),
    ("Weyregg am Attersee", 1600, 100000, ""),
    ("Unterach am Attersee", 1550, 75000, ""),
    ("Nußdorf am Attersee", 1200, 70000, ""),
    ("Steinbach am Attersee", 850, 65000, ""),
]
g0 = 10
gl = g0 + len(gem) - 1
gs = gl + 1
avg_ue = f"AVERAGE(GuV_10J!$B${G_ROWS['g_ue']}:$K${G_ROWS['g_ue']})"
for i, (n, ew, na, x) in enumerate(gem):
    rr = g0 + i
    P.cell(row=rr, column=1, value=n).font = BLACK
    put(P, f"B{rr}", ew, BLUE, NUM, YELLOW)
    put(P, f"C{rr}", na, BLUE, NUM, YELLOW)
    put(P, f"D{rr}", x, BLUE)
    P[f"D{rr}"].alignment = Alignment(horizontal="center")
    put(P, f"E{rr}", f"=(1-$D$5-$D$6)*($D$7*B{rr}/$B${gs}+(1-$D$7)*C{rr}/$C${gs})+IF(D{rr}=\"x\",$D$6,0)", BLACK, PCT)
    put(P, f"F{rr}", f"=$D$4*E{rr}", BLACK, EUR)
    put(P, f"G{rr}", f"={avg_ue}*E{rr}/(1-$D$5)", BLACK, EUR)
    put(P, f"H{rr}", f"=F{rr}/B{rr}", BLACK, EUR)
put(P, f"A{gs}", "Summe Gemeinden", bold=True)
put(P, f"B{gs}", f"=SUM(B{g0}:B{gl})", BLACK, NUM, SUMFILL, bold=True)
put(P, f"C{gs}", f"=SUM(C{g0}:C{gl})", BLACK, NUM, SUMFILL, bold=True)
put(P, f"E{gs}", f"=SUM(E{g0}:E{gl})", BLACK, PCT, SUMFILL, bold=True)
put(P, f"F{gs}", f"=SUM(F{g0}:F{gl})", BLACK, EUR, SUMFILL, bold=True)
put(P, f"G{gs}", f"=SUM(G{g0}:G{gl})", BLACK, EUR, SUMFILL, bold=True)
put(P, f"H{gs}", f"=F{gs}/B{gs}", BLACK, EUR, SUMFILL, bold=True)
rt = gs + 1
put(P, f"A{rt}", "Tourismusverband Attersee-Attergau")
put(P, f"E{rt}", "=D5", BLACK, PCT)
put(P, f"F{rt}", f"=$D$4*E{rt}", BLACK, EUR)
rk = rt + 1
put(P, f"A{rk}", "Kontrolle: Summe = zu verteilende Investition", bold=True)
put(P, f"E{rk}", f"=E{gs}+E{rt}", BLACK, PCT, bold=True)
put(P, f"F{rk}", f"=F{gs}+F{rt}", BLACK, EUR, bold=True)
put(P, f"G{rk}", f'=IF(ABS(F{rk}-D4)<1,"OK","FEHLER")', BLACK, bold=True)
P[f"A{rk+2}"] = ("Pachtrückfluss: Ø-Überschuss der Gemeinde-GmbH (GuV_10J) nach Gemeinde-Anteil verteilt; "
                 "TVB-Anteil ohne Rückfluss angenommen (Zuschuss).")
P[f"A{rk+2}"].font = Font(name=F, size=9, italic=True)

# ---------------------------------------------------------------- Szenarien
S = wb.create_sheet("Szenarien")
S["A1"] = "Szenarienvergleich – Normaljahr"
S["A1"].font = TITLE
hdr(S, 3, ["Kennzahl", "Einheit", "pessimistisch", "Basis", "optimistisch"], [52, 10, 17, 17, 17])
srows = [
    ("Besuche pro Jahr", "Besuche", lambda c: f"={u('bes', c)}", NUM),
    ("Ø Besucher/Tag Hauptsaison", "Besucher", lambda c: f"={a('g_hs', c)}", NUM),
    ("Umsatz netto", "€", lambda c: f"={u('u_sum', c)}", EUR),
    ("Umsatz je Besuch", "€", lambda c: f"={u('u_kopf', c)}", DEC),
    ("Betriebskosten vor Pacht", "€", lambda c: f"={b('k_sum', c)}", EUR),
    ("Pacht an Gemeinde-GmbH", "€", lambda c: f"={b('pacht', c)}", EUR),
    ("Betreiberergebnis (EBITDA)", "€", lambda c: f"={b('ebitda', c)}", EUR),
    ("EBITDA-Marge", "%", lambda c: f"={b('marge', c)}", PCT),
    ("Break-even Besuche (Betreiber)", "Besuche", lambda c: f"={b('be', c)}", NUM),
    ("Break-even in % der Planbesuche", "%", lambda c: f"={b('be_q', c)}", PCT),
    ("Überschuss Gemeinde-GmbH p. a. (Pacht – Rücklage – Verwaltung)", "€",
     lambda c: f"={b('pacht', c)}-{INV_BAU}*{a('rl', c)}-{a('g_verw', c)}", EUR),
    ("Gemeinde-Investition (netto nach Förderung)", "€", lambda c: f"={INV_GEM}", EUR),
    ("Rückflussdauer aus Pacht (statisch)", "Jahre", lambda c: f"=IF(C14>0,{INV_GEM}/C14,\"n/a\")".replace("C14", f"{c}14"), YRS),
    ("Tragbarer Pachtsatz (bei Ziel-Marge Betreiber)", "%",
     lambda c: f"=MAX(0,{b('ebitda_v', c)}/{u('u_sum', c)}-{a('ziel_m', c)})", PCT),
    ("Tragbare Pacht p. a.", "€", lambda c: f"={c}17*{u('u_sum', c)}", EUR),
]
for i, (lab, unit, fn, fmt) in enumerate(srows):
    rr = 4 + i
    S.cell(row=rr, column=1, value=lab).font = BLACK
    S.cell(row=rr, column=2, value=unit).font = BLACK
    for c in SC:
        put(S, f"{c}{rr}", fn(c), GREEN if fn(c).count("!") == 1 and "-" not in fn(c) and "/" not in fn(c) else BLACK, fmt)
S["A21"] = ("Lesart: Die Pacht deckt Instandhaltung und einen Teil der Investition (Teil-Tilgung). Eine Vollrefinanzierung "
            "über die Pacht ist nicht vorgesehen – der Rest ist Standort- und Tourismusinfrastruktur der Gemeinden "
            "(Umwegrentabilität: Kommunalsteuer, Nächtigungen, Ganzjahresangebot).")
S["A21"].font = Font(name=F, italic=True, size=9)
S["A21"].alignment = Alignment(wrap_text=True, vertical="top")
S.merge_cells("A21:E21")
S.row_dimensions[21].height = 48

# order sheets
wb._sheets = [S, A, RP, I, U, B, SP, G, P]
wb.active = 0
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
